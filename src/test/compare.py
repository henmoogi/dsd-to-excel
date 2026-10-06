# Python판(o??.xlsx)과 HTML판(j??.xlsx) 결과를 셀 단위로 비교
import sys, os, glob, collections
from openpyxl import load_workbook
pdir, jdir = sys.argv[1], sys.argv[2]
tot = collections.Counter(); samples = collections.defaultdict(list)
def note(kind, msg):
    tot[kind] += 1
    if len(samples[kind]) < 4: samples[kind].append(msg)
def font_sig(c):
    f = c.font; col = f.color.rgb[-6:] if f and f.color is not None and isinstance(f.color.rgb, str) else None
    return (f.name, float(f.sz or 0), bool(f.b), bool(f.u), col)
def al_sig(c):
    a = c.alignment; return (a.horizontal, a.vertical, bool(a.wrap_text))
def fill_sig(c):
    fl = c.fill; return (fl.fill_type, fl.fgColor.rgb[-6:] if fl.fill_type else None)
for i in range(len(glob.glob(os.path.join(jdir, "j*.xlsx")))):
    p = load_workbook(os.path.join(pdir, f"o{i:02d}.xlsx")); j = load_workbook(os.path.join(jdir, f"j{i:02d}.xlsx"))
    tag = f"#{i:02d}"
    jn = [n for n in j.sheetnames if n != "검증"]   # 검증 시트는 HTML판에만 있다
    if p.sheetnames != jn: note("시트목록", f"{tag} {p.sheetnames} / {jn}"); continue
    if p["목차"]["A2"].value != j["목차"]["A2"].value: note("원본이름", f"{tag}")
    for name in p.sheetnames:
        a, b = p[name], j[name]
        cells = set()
        for ws, s in ((a, 'p'), (b, 'j')):
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None: cells.add(c.coordinate)
        for co in sorted(cells):
            if name == "목차" and co in ("A3", "A4"): continue   # 변환 시각, 검증 요약 링크
            ca, cb = a[co], b[co]
            va, vb = ca.value, cb.value
            if isinstance(va, (int, float)) and isinstance(vb, (int, float)):
                if abs(va - vb) > 1e-9: note("값", f"{tag} {name}!{co} {va!r} / {vb!r}")
            elif va != vb: note("값", f"{tag} {name}!{co} {va!r} / {vb!r}")
            if ca.number_format != cb.number_format: note("표시형식", f"{tag} {name}!{co} {ca.number_format} / {cb.number_format}")
            if font_sig(ca) != font_sig(cb): note("글꼴", f"{tag} {name}!{co} {font_sig(ca)} / {font_sig(cb)}")
            if al_sig(ca) != al_sig(cb): note("정렬", f"{tag} {name}!{co} {al_sig(ca)} / {al_sig(cb)}")
            if fill_sig(ca) != fill_sig(cb) and fill_sig(cb)[1] != "FFC7CE": note("채우기", f"{tag} {name}!{co} {fill_sig(ca)} / {fill_sig(cb)}")
            ha, hb = ca.hyperlink, cb.hyperlink
            la = ha.location if ha else None; lb = hb.location if hb else None
            if la != lb: note("링크", f"{tag} {name}!{co} {la} / {lb}")
            if hb is not None and hb.target: note("외부링크남음", f"{tag} {name}!{co} {hb.target}")
        ba = {c.coordinate for row in a.iter_rows() for c in row if c.border.left.style}
        bb = {c.coordinate for row in b.iter_rows() for c in row if c.border.left.style}
        if ba != bb: note("테두리", f"{tag} {name} 차이 {len(ba ^ bb)}칸 예 {sorted(ba ^ bb)[:3]}")
        ma = {str(m) for m in a.merged_cells.ranges}; mb = {str(m) for m in b.merged_cells.ranges}
        if ma != mb: note("병합", f"{tag} {name} {sorted(ma - mb)[:3]} / {sorted(mb - ma)[:3]}")
        if a.freeze_panes != b.freeze_panes: note("틀고정", f"{tag} {name} {a.freeze_panes} / {b.freeze_panes}")
        oa = {r for r, d in a.row_dimensions.items() if d.outline_level}; ob = {r for r, d in b.row_dimensions.items() if d.outline_level}
        if oa != ob: note("그룹", f"{tag} {name} {len(oa)} / {len(ob)} 예 {sorted(oa ^ ob)[:5]}")
        ra = {r: d.height for r, d in a.row_dimensions.items() if d.height}; rb = {r: d.height for r, d in b.row_dimensions.items() if d.height}
        if ra != rb: note("행높이", f"{tag} {name} 예 {[(r, ra.get(r), rb.get(r)) for r in sorted(set(ra) ^ set(rb) | {r for r in ra if r in rb and ra[r] != rb[r]})][:3]}")
        def widths(ws):
            out = {}
            for key, d in ws.column_dimensions.items():
                if d.width:
                    lo = d.min or 0; hi = d.max or 0
                    from openpyxl.utils import column_index_from_string
                    k = column_index_from_string(key)
                    for ci in range(min(lo or k, k), max(hi or k, k) + 1): out[ci] = d.width
            return out
        wa, wb2 = widths(a), widths(b)
        for ci in sorted(set(wa) | set(wb2)):
            x, y = wa.get(ci), wb2.get(ci)
            if x is None or y is None or abs(x - y) > 0.15: note("열너비", f"{tag} {name} 열{ci} {x} / {y}")
        if a.sheet_view.showGridLines != b.sheet_view.showGridLines: note("눈금선", f"{tag} {name}")
        sa = a.sheet_properties.outlinePr.summaryBelow if a.sheet_properties.outlinePr else None
        sb = b.sheet_properties.outlinePr.summaryBelow if b.sheet_properties.outlinePr else None
        if bool(oa) and sa != sb: note("요약행위치", f"{tag} {name} {sa} / {sb}")
    tot["파일"] += 1
print("비교한 파일:", tot.pop("파일", 0))
print("차이 없음" if not tot else "차이:")
for k, n in tot.most_common():
    print(f"  {k}: {n}건"); [print("     ", s) for s in samples[k]]
