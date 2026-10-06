import sys, glob, os, collections
from openpyxl import load_workbook
d = sys.argv[1]; show = sys.argv[2] if len(sys.argv) > 2 else "불일치"
tot = collections.Counter(); kinds = collections.Counter()
for f in sorted(glob.glob(os.path.join(d, "j*.xlsx"))):
    wb = load_workbook(f); ws = wb["검증"]
    toc = wb["목차"]["A2"].value[4:30]
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not row[0]: continue
        tot[row[0]] += 1; kinds[(row[0], row[1])] += 1
        if row[0] in show.split(","):
            print(f"{os.path.basename(f)[:3]} {row[0]} {row[1]} {row[2]} | {str(row[3])[:28]} | {str(row[4])[:14]} | {row[5]} | {row[6]} | {row[7]} | {str(row[8])[:70]}")
print(dict(tot)); print(sorted(kinds.items()))
